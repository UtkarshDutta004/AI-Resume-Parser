"""Generate the PDF and DOCX sample resumes omitted from version control."""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"
RESUME_1 = [
    "Jordan Lee",
    "jordan.lee@example.com | +1 (212) 555-0198",
    "https://www.linkedin.com/in/jordan-lee | https://github.com/jordanlee",
    "",
    "SUMMARY",
    "Data analyst focused on practical machine learning products.",
    "",
    "SKILLS",
    "Python, SQL, Machine Learning, AWS, Git",
    "",
    "EXPERIENCE",
    "Data Analyst, Northstar Analytics | 2020 - 2024",
    "Created metrics and predictive models for product teams.",
    "",
    "EDUCATION",
    "Master of Science, Columbia University | 2020",
    "",
    "PROJECTS",
    "Customer Churn Model - Developed a churn prediction workflow.",
    "",
    "CERTIFICATIONS",
    "AWS Certified Solutions Architect",
]
RESUME_2 = [
    "Alex Morgan",
    "alex.morgan@example.org | (312) 555-0177",
    "https://www.linkedin.com/in/alex-morgan",
    "",
    "PROFESSIONAL EXPERIENCE",
    "Frontend Developer, Pixel Works | 2019 - 2023",
    "Delivered responsive product experiences with React and TypeScript.",
    "",
    "TECHNICAL SKILLS",
    "JavaScript, TypeScript, React, HTML, CSS, Node.js",
    "",
    "EDUCATION",
    "Bachelor of Arts, Lakeside College | 2019",
    "",
    "PROJECTS",
    "Design System - Created reusable interface components.",
    "",
    "CERTIFICATIONS",
    "Google UX Design Certificate",
]


def create_pdf(path: Path, lines: list[str]) -> None:
    """Write a readable, single-page text PDF."""
    pdf = canvas.Canvas(str(path), pagesize=letter)
    text = pdf.beginText(54, 740)
    text.setFont("Helvetica", 10)
    text.setLeading(15)
    for line in lines:
        text.textLine(line)
    pdf.drawText(text)
    pdf.save()


def create_docx(path: Path, lines: list[str]) -> None:
    """Write a DOCX sample, using heading styles for section labels."""
    document = Document()
    for line in lines:
        if line and line.upper() in {
            "SUMMARY", "PROFESSIONAL EXPERIENCE", "TECHNICAL SKILLS",
            "EDUCATION", "PROJECTS", "CERTIFICATIONS",
        }:
            document.add_heading(line.title(), level=2)
        else:
            document.add_paragraph(line)
    document.save(path)


def main() -> None:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    create_pdf(SAMPLES / "sample_resume_1.pdf", RESUME_1)
    create_docx(SAMPLES / "sample_resume_2.docx", RESUME_2)
    print("Generated sample_resume_1.pdf and sample_resume_2.docx in samples/.")


if __name__ == "__main__":
    main()
