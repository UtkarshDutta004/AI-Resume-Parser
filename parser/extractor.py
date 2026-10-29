"""Text extraction from supported resume formats; uploads are handled in memory."""

from __future__ import annotations

from io import BytesIO
from zipfile import BadZipFile


class ExtractionError(ValueError):
    """Raised when a supported file cannot be decoded into text."""


def extract_text(file_bytes: bytes, extension: str) -> str:
    """Extract text from PDF, DOCX, or UTF-8 text bytes without persisting uploads."""
    if extension == "txt":
        try:
            return file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ExtractionError("Unable to read this TXT file. Use UTF-8 encoded text.") from exc
    if extension == "pdf":
        try:
            from pdfminer.high_level import extract_text as pdf_extract_text
            from pdfminer.pdfparser import PDFException
        except ImportError as exc:
            raise ExtractionError("PDF support is unavailable; install the declared dependencies.") from exc
        try:
            return pdf_extract_text(BytesIO(file_bytes))
        except (OSError, ValueError, PDFException) as exc:
            raise ExtractionError("Unable to read this PDF file.") from exc
    if extension == "docx":
        try:
            from docx import Document
            from docx.opc.exceptions import PackageNotFoundError
        except ImportError as exc:
            raise ExtractionError("DOCX support is unavailable; install the declared dependencies.") from exc
        try:
            document = Document(BytesIO(file_bytes))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
            return "\n".join(parts)
        except (OSError, ValueError, BadZipFile, PackageNotFoundError) as exc:
            raise ExtractionError("Unable to read this DOCX file.") from exc
    raise ExtractionError("Unsupported file type.")
